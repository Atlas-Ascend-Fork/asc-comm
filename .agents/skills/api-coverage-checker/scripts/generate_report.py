"""
Compare code APIs against documented APIs, generate xlsx + md reports.

Usage:
    python generate_report.py --code <code_apis.json> --docs <doc_apis.json>
        --out-xlsx <output.xlsx> --out-md <output.md>
"""
import argparse
import json
import math
import os
import sys
from collections import defaultdict

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

MODULE_DISPLAY = {
    "ccu": "CCU API (ccu)",
}

CATEGORY_DISPLAY = {
    "resource_allocation_operation": "资源创建与操作",
    "arg_load_store": "参数加载/存储",
    "data_movement": "数据搬运",
    "synchronization": "同步",
    "execution_control": "流程控制",
    "host_launch": "Host注册/启动",
    "ccu_support": "CCU支撑类型",
    "vector_compute": "矢量计算",
    "cube_compute": "矩阵计算",
    "data_move": "数据搬运",
    "sync": "同步控制",
    "scalar": "标量计算",
    "sys_var": "系统变量",
    "activation": "激活函数",
    "math": "数学运算",
    "matmul": "矩阵乘",
    "conv": "卷积",
    "reduce": "规约操作",
    "normalization": "归一化",
    "quantization": "量化",
    "hccl": "集合通信",
    "sort": "排序",
    "pad": "填充",
    "transpose": "转置",
    "filter": "过滤/掩码",
    "select": "选择",
    "index": "索引",
    "tensor": "Tensor",
    "utils": "工具/调试",
    "other": "其他",
}


def load_json(path):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def compare(code_data, doc_data):
    """Compare code APIs against documented APIs and identify gaps."""
    doc_names = set()
    doc_by_name = {}
    for api in doc_data.get("apis", []):
        name = api["api_name"].lower().strip()
        doc_names.add(name)
        doc_by_name[name] = api

    code_apis = code_data.get("apis", [])

    # Build results
    covered = []
    missing = []
    for api in code_apis:
        name = api["api_name"].lower().strip()
        doc_api = doc_by_name.get(name)
        if doc_api:
            api["doc_file"] = doc_api.get("doc_file", "")
            api["doc_npu_arch"] = doc_api.get("npu_arch")
            api["doc_signature"] = doc_api.get("signature")
            covered.append(api)
        else:
            missing.append(api)

    # Module-level stats
    module_stats = {}
    for api in code_apis:
        mod = api["module"]
        module_stats.setdefault(mod, {"total": 0, "covered": 0, "missing": 0})
        module_stats[mod]["total"] += 1
    for api in covered:
        module_stats[api["module"]]["covered"] += 1
    for api in missing:
        module_stats[api["module"]]["missing"] += 1

    # Category-level stats for missing
    cat_stats = defaultdict(int)
    for api in missing:
        cat_stats[api["category"]] += 1

    return {
        "covered": covered,
        "missing": missing,
        "module_stats": module_stats,
        "cat_stats": dict(cat_stats),
        "total_code": len(code_apis),
        "total_docs": len(doc_data.get("apis", [])),
        "total_covered": len(covered),
        "total_missing": len(missing),
        "coverage_rate": round(len(covered) / len(code_apis) * 100, 1) if code_apis else 0,
    }


def generate_xlsx(result, output_path):
    """Generate formatted xlsx workbook."""
    if not HAS_OPENPYXL:
        print("openpyxl not installed, skipping xlsx generation")
        return

    wb = openpyxl.Workbook()
    hdr_font = Font(bold=True, size=11, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="4472C4")
    hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_align = Alignment(vertical="top", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )

    def style_header(ws, row, cols):
        for col in range(1, cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = hdr_align
            cell.border = thin_border

    # Sheet 1: Summary
    ws1 = wb.active
    ws1.title = "缺口总览"
    headers1 = ["模块", "代码接口总数", "已文档化", "缺失数", "覆盖率"]
    ws1.append(headers1)
    style_header(ws1, 1, len(headers1))

    for mod in MODULE_DISPLAY:
        stats = result["module_stats"].get(mod, {})
        total = stats.get("total", 0)
        covered_cnt = stats.get("covered", 0)
        missing_cnt = stats.get("missing", 0)
        rate = f"{covered_cnt / total * 100:.1f}%" if total > 0 else "N/A"
        ws1.append([MODULE_DISPLAY[mod], total, covered_cnt, missing_cnt, rate])

    # Total row
    ws1.append(["合计", result["total_code"], result["total_covered"],
                result["total_missing"], f"{result['coverage_rate']}%"])
    for col in range(1, 6):
        ws1.cell(row=ws1.max_row, column=col).font = Font(bold=True)

    for col in range(1, 6):
        ws1.column_dimensions[get_column_letter(col)].width = 18
    ws1.freeze_panes = "A2"

    # Sheet 2: Missing API List
    ws2 = wb.create_sheet("缺失接口清单")
    headers2 = ["序号", "API名称", "模块", "功能类别", "接口签名", "源文件", "代码ARCH", "备注"]
    ws2.append(headers2)
    style_header(ws2, 1, len(headers2))

    for i, api in enumerate(result["missing"], 1):
        mod_disp = MODULE_DISPLAY.get(api["module"], api["module"])
        cat_disp = CATEGORY_DISPLAY.get(api["category"], api["category"])
        sig = api.get("signature", "") or ""
        file_src = api.get("file", "")
        code_arch = api.get("npu_arch", "") or "全架构"
        ws2.append([i, api["api_name"], mod_disp, cat_disp, sig, file_src, code_arch, ""])

    col_widths2 = [6, 28, 22, 14, 50, 50, 16, 20]
    for col, w in enumerate(col_widths2, 1):
        ws2.column_dimensions[get_column_letter(col)].width = w
    ws2.freeze_panes = "A2"

    # Sheet 3: Category-level stats
    ws3 = wb.create_sheet("按功能类别统计")
    headers3 = ["功能类别", "缺失数", "占比"]
    ws3.append(headers3)
    style_header(ws3, 1, len(headers3))

    sorted_cats = sorted(result["cat_stats"].items(), key=lambda x: -x[1])
    total_missing = result["total_missing"]
    for cat, count in sorted_cats:
        cat_disp = CATEGORY_DISPLAY.get(cat, cat)
        ratio = f"{count / total_missing * 100:.1f}%" if total_missing > 0 else "0%"
        ws3.append([cat_disp, count, ratio])

    ws3.column_dimensions["A"].width = 18
    ws3.column_dimensions["B"].width = 12
    ws3.column_dimensions["C"].width = 10
    ws3.freeze_panes = "A2"

    wb.save(output_path)
    print(f"xlsx saved to {output_path}")


def generate_md(result, output_path):
    """Generate markdown analysis report."""
    lines = []
    lines.append("# API 覆盖度分析报告")
    lines.append("")
    lines.append(f"- 代码接口总数：{result['total_code']}")
    lines.append(f"- 已文档化接口：{result['total_covered']}")
    lines.append(f"- 缺失接口：{result['total_missing']}")
    lines.append(f"- 总体覆盖率：**{result['coverage_rate']}%**")
    lines.append("")

    # Module summary table
    lines.append("## 一、按模块统计")
    lines.append("")
    lines.append("| 模块 | 代码接口 | 已文档化 | 缺失 | 覆盖率 |")
    lines.append("|------|---------|---------|------|--------|")

    for mod in MODULE_DISPLAY:
        stats = result["module_stats"].get(mod, {})
        total = stats.get("total", 0)
        covered_cnt = stats.get("covered", 0)
        missing_cnt = stats.get("missing", 0)
        rate = f"{covered_cnt / total * 100:.1f}%" if total > 0 else "-"
        if total > 0:
            lines.append(f"| {MODULE_DISPLAY[mod]} | {total} | {covered_cnt} | {missing_cnt} | {rate} |")

    lines.append("")

    # Category stats
    lines.append("## 二、按功能类别统计（缺失部分）")
    lines.append("")
    lines.append("| 功能类别 | 缺失数 | 占比 |")
    lines.append("|---------|-------|------|")
    sorted_cats = sorted(result["cat_stats"].items(), key=lambda x: -x[1])
    for cat, count in sorted_cats:
        cat_disp = CATEGORY_DISPLAY.get(cat, cat)
        ratio = f"{count / result['total_missing'] * 100:.1f}%" if result["total_missing"] > 0 else "0%"
        lines.append(f"| {cat_disp} | {count} | {ratio} |")
    lines.append("")

    # Top missing APIs
    lines.append("## 三、缺失接口示例（前 30）")
    lines.append("")
    lines.append("| API名称 | 模块 | 功能类别 | 源文件 |")
    lines.append("|---------|------|---------|--------|")
    for api in result["missing"][:30]:
        mod_disp = MODULE_DISPLAY.get(api["module"], api["module"])
        cat_disp = CATEGORY_DISPLAY.get(api["category"], api["category"])
        fname = api.get("file", "").split("/")[-1] if api.get("file") else ""
        lines.append(f"| `{api['api_name']}` | {mod_disp} | {cat_disp} | {fname} |")
    lines.append("")

    # Coverage bar chart
    lines.append("## 四、覆盖率可视化")
    lines.append("")
    rate = result["coverage_rate"]
    bar_len = 40
    filled = int(bar_len * rate / 100)
    bar = "█" * filled + "░" * (bar_len - filled)
    lines.append(f"```")
    lines.append(f"总体: {bar} {rate}%")
    lines.append("")

    for mod in MODULE_DISPLAY:
        stats = result["module_stats"].get(mod, {})
        total = stats.get("total", 0)
        covered_cnt = stats.get("covered", 0)
        if total > 0:
            r = covered_cnt / total * 100
            f = int(bar_len * r / 100)
            b = "█" * f + "░" * (bar_len - f)
            lines.append(f"{MODULE_DISPLAY[mod]:20s}: {b} {r:.1f}%")
    lines.append("```")

    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"md saved to {output_path}")


def main():
    parser = argparse.ArgumentParser(description="Generate API coverage gap report")
    parser.add_argument("--code", required=True, help="JSON file from extract_code_apis.py")
    parser.add_argument("--docs", required=True, help="JSON file from extract_doc_apis.py")
    parser.add_argument("--out-xlsx", required=True)
    parser.add_argument("--out-md", required=True)
    args = parser.parse_args()

    code_data = load_json(args.code)
    doc_data = load_json(args.docs)
    result = compare(code_data, doc_data)

    print(f"Code APIs: {result['total_code']}")
    print(f"Doc APIs:  {result['total_docs']}")
    print(f"Covered:   {result['total_covered']}")
    print(f"Missing:   {result['total_missing']}")
    print(f"Coverage:  {result['coverage_rate']}%")

    generate_xlsx(result, args.out_xlsx)
    generate_md(result, args.out_md)


if __name__ == "__main__":
    main()
